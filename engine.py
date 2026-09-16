import io, re, unicodedata
from typing import Callable, List, Tuple
import fitz
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

COLS=['No.','CIV INICIO','CIV FIN','DIRECCIÓN DE LA OBRA INICIO','DIRECCIÓN DE LA OBRA FIN','CONTRATISTA','FECHA INICIO','FECHA FIN','HORARIO DE TRABAJO','HORARIO DE CIERRE','No. CONTRATO','OBSERVACIONES','AUTORIZADO','LOCALIDAD','ING. RESPONSABLE','No RADICADO SDM']

def clean(s):
    return re.sub(r'\s+',' ',str(s or '').replace('\u00a0',' ')).strip()

def norm(s):
    s=clean(s).upper()
    s=''.join(c for c in unicodedata.normalize('NFD',s) if unicodedata.category(c)!='Mn')
    return re.sub(r'[^A-Z0-9]+','',s)

def contract_canonical(s):
    s=clean(s).upper().replace(' ','')
    m=re.search(r'(?:SDM-)?(\d{3,5})-(20\d{2})$',s)
    if m: return f'{m.group(2)}-{m.group(1)}'
    m=re.search(r'(20\d{2})-(\d{3,5})$',s)
    if m: return f'{m.group(1)}-{m.group(2)}'
    return s

def is_civ(s):
    return bool(re.fullmatch(r'\d{5,10}',clean(s)))

def find_verticals(page, y0=250):
    xs=[]
    for dr in page.get_drawings():
        for item in dr.get('items',[]):
            if item[0]=='l':
                a,b=item[1],item[2]
                if abs(a.x-b.x)<1.5 and max(a.y,b.y)-min(a.y,b.y)>100 and max(a.y,b.y)>y0:
                    xs.append(round(a.x,1))
    xs=sorted(xs); out=[]
    for x in xs:
        if not out or abs(x-out[-1])>2: out.append(x)
    return out

def find_horizontal_rows(page):
    lines=[]
    for dr in page.get_drawings():
        for item in dr.get('items',[]):
            if item[0]=='l':
                a,b=item[1],item[2]
                if abs(a.y-b.y)<1.5 and (b.x-a.x)>1700:
                    lines.append(round((a.y+b.y)/2,1))
    lines=sorted(lines); out=[]
    for y in lines:
        if not out or abs(y-out[-1])>2: out.append(y)
    return out

def extract_page(page):
    words=page.get_text('words')
    xs=find_verticals(page)
    hfull=[]
    for dr in page.get_drawings():
        for item in dr.get('items',[]):
            if item[0]=='l':
                a,b=item[1],item[2]
                if abs(a.y-b.y)<1.5 and (b.x-a.x)>1700:
                    hfull.append((min(a.x,b.x),max(a.x,b.x),round((a.y+b.y)/2,1)))
    if hfull:
        xmin=min(x[0] for x in hfull); xmax=max(x[1] for x in hfull)
        xs=[xmin]+xs+[xmax]
    xs=sorted(set(round(x,1) for x in xs))
    if len(xs)>17:
        target=[74,123,191,259,358,456,585,652,718,796,874,952,1593,1688,1785,1886,2009]
        chosen=[]
        for t in target:
            nearest=min(xs,key=lambda x:abs(x-t))
            if nearest not in chosen: chosen.append(nearest)
        if len(chosen)==17: xs=chosen
    if len(xs)!=17: return []
    ys=find_horizontal_rows(page)
    if len(ys)<2: return []
    out=[]
    for ya,yb in zip(ys,ys[1:]):
        if yb-ya<15: continue
        rowwords=[w for w in words if w[1]>=ya-1 and w[3]<=yb+1]
        nums=[w for w in rowwords if xs[0]-3<=w[0]<=xs[1]+3 and re.fullmatch(r'\d{3,8}',w[4].strip())]
        if not nums: continue
        cells=['' for _ in range(16)]
        for w in rowwords:
            x0,y0,x1,y1,text,*_=w; cx=(x0+x1)/2; j=None
            for k in range(16):
                if xs[k]-1<=cx<=xs[k+1]+1:
                    j=k; break
            if j is not None: cells[j]=(cells[j]+' '+text).strip()
        cells=[clean(c) for c in cells]
        if is_civ(cells[1]) or (cells[0].isdigit() and cells[0]):
            for j in [1,2]:
                m=re.match(r'(\d{5,10})',cells[j])
                cells[j]=m.group(1) if m else (cells[j] if cells[j] in ['N/A','-'] else cells[j])
            out.append(cells)
    return out

def extract_pdf(data:bytes, progress:Callable[[int,str],None]|None=None):
    doc=fitz.open(stream=data,filetype='pdf')
    records=[]; page_diag=[]; total=len(doc)
    for pi,page in enumerate(doc):
        rows=extract_page(page)
        txt=page.get_text('text').upper()
        section='SECCIÓN 1' if 'SECCIÓN 1.' in txt else ('SECCIÓN 2' if 'SECCIÓN 2.' in txt else '')
        page_diag.append((pi+1,len(rows)))
        for r in rows:
            r=r[:16]+['']*max(0,16-len(r))
            rec=dict(zip(COLS,r[:16])); rec['PÁGINA PDF']=pi+1; rec['SECCIÓN']=section
            rec['CONTRATO CANÓNICO']=contract_canonical(rec['No. CONTRATO'])
            auth=norm(rec['AUTORIZADO']); obs=norm(rec['OBSERVACIONES'])
            if 'EMERGENCIA' in auth or 'FORMALIZACIONDEEMERGENCIA' in auth:
                state='FORMALIZACIÓN DE EMERGENCIA'
            elif auth in ('NO','NOAUTORIZADO','NOAUTORIZADA'):
                state='NO AUTORIZADO'
            elif auth in ('SI','SIAUTORIZADO','AUTORIZADO','AUTORIZADOCONOBSERVACIONES') or 'AUTORIZA' in obs:
                state='AUTORIZADO'
            else:
                state=clean(rec['AUTORIZADO']) or 'OTRO'
            rec['ESTADO INTERPRETADO']=state
            records.append(rec)
        if progress:
            pct=int(((pi+1)/max(total,1))*100)
            progress(pct,f'Analizando página {pi+1:,} de {total:,} · {len(records):,} registros encontrados')
    doc.close()
    return pd.DataFrame(records),page_diag

def contract_variants(contract):
    k=contract_canonical(contract); variants={k,norm(contract)}
    for value in [k]:
        m=re.match(r'(20\d{2})-(\d{3,5})$',value or '')
        if m:
            y,n=m.group(1),m.group(2)
            variants.update({norm(f'SDM-{n}-{y}'),norm(f'SDM-{n}- {y}'),norm(f'{n}-{y}'),norm(f'{y}-{n}'),norm(f'{n}{y}'),norm(f'{y}{n}')})
    return {v for v in variants if v}

def source_validation(data,df):
    doc=fitz.open(stream=data,filetype='pdf'); full_norm=norm('\n'.join((p.get_text('text') or '') for p in doc)); doc.close()
    g=df.copy()
    if 'CONTRATO CANÓNICO' not in g.columns: g['CONTRATO CANÓNICO']=g['No. CONTRATO'].map(contract_canonical)
    g['CONTRATISTA']=g['CONTRATISTA'].fillna('').map(clean); g['CONTRATO CANÓNICO']=g['CONTRATO CANÓNICO'].fillna('').map(clean)
    g=g.groupby(['CONTRATISTA','CONTRATO CANÓNICO'],dropna=False).size().reset_index(name='REGISTROS_EXTRAÍDOS')
    rows=[]
    for _,r in g.iterrows():
        company=clean(r['CONTRATISTA']); contract=clean(r['CONTRATO CANÓNICO']); c=norm(company)
        mentions_company=full_norm.count(c) if c else 0
        mentions_contract=max([full_norm.count(norm(v)) for v in contract_variants(contract)] or [0])
        extracted=int(r['REGISTROS_EXTRAÍDOS'])
        rows.append({'CONTRATISTA':company,'CONTRATO CANÓNICO':contract,'MENCIONES CONTRATISTA EN PDF':mentions_company,'MENCIONES CONTRATO EN PDF':mentions_contract,'REGISTROS EXTRAÍDOS':extracted,'DIFERENCIA REFERENCIAL':mentions_company-extracted,'VALIDACIÓN':'⚠️ REVISAR' if mentions_company!=extracted else '✅ OK'})
    return pd.DataFrame(rows).sort_values(['VALIDACIÓN','REGISTROS EXTRAÍDOS'],ascending=[True,False]).reset_index(drop=True) if rows else pd.DataFrame()

def make_excel(df, validation=None):
    bio=io.BytesIO()
    with pd.ExcelWriter(bio,engine='openpyxl') as writer:
        df.to_excel(writer,index=False,sheet_name='COI completo')
        summary=pd.DataFrame([{'Indicador':'Registros','Valor':len(df)},{'Indicador':'Autorizados','Valor':int((df['ESTADO INTERPRETADO']=='AUTORIZADO').sum())},{'Indicador':'No autorizados','Valor':int((df['ESTADO INTERPRETADO']=='NO AUTORIZADO').sum())},{'Indicador':'Emergencias','Valor':int((df['ESTADO INTERPRETADO']=='FORMALIZACIÓN DE EMERGENCIA').sum())},{'Indicador':'24 HORAS','Valor':int(df['HORARIO DE TRABAJO'].map(norm).eq('24HORAS').sum())},{'Indicador':'Empresas','Valor':df['CONTRATISTA'].nunique()},{'Indicador':'Contratos','Valor':df['CONTRATO CANÓNICO'].nunique()}])
        summary.to_excel(writer,index=False,sheet_name='Resumen')
        emp=df.groupby('CONTRATISTA',dropna=False).agg(REGISTROS=('No.','count'),AUTORIZADOS=('ESTADO INTERPRETADO',lambda s:(s=='AUTORIZADO').sum()),NO_AUTORIZADOS=('ESTADO INTERPRETADO',lambda s:(s=='NO AUTORIZADO').sum()),EMERGENCIAS=('ESTADO INTERPRETADO',lambda s:(s=='FORMALIZACIÓN DE EMERGENCIA').sum()),CONTRATOS=('CONTRATO CANÓNICO','nunique')).reset_index().sort_values('REGISTROS',ascending=False)
        emp.to_excel(writer,index=False,sheet_name='Empresas')
        con=df.groupby(['CONTRATO CANÓNICO','No. CONTRATO','CONTRATISTA'],dropna=False).agg(REGISTROS=('No.','count'),AUTORIZADOS=('ESTADO INTERPRETADO',lambda s:(s=='AUTORIZADO').sum()),NO_AUTORIZADOS=('ESTADO INTERPRETADO',lambda s:(s=='NO AUTORIZADO').sum()),EMERGENCIAS=('ESTADO INTERPRETADO',lambda s:(s=='FORMALIZACIÓN DE EMERGENCIA').sum())).reset_index().sort_values('REGISTROS',ascending=False)
        con.to_excel(writer,index=False,sheet_name='Contratos')
        hor=df.groupby('HORARIO DE TRABAJO',dropna=False).agg(REGISTROS=('No.','count'),AUTORIZADOS=('ESTADO INTERPRETADO',lambda s:(s=='AUTORIZADO').sum()),NO_AUTORIZADOS=('ESTADO INTERPRETADO',lambda s:(s=='NO AUTORIZADO').sum()),EMERGENCIAS=('ESTADO INTERPRETADO',lambda s:(s=='FORMALIZACIÓN DE EMERGENCIA').sum())).reset_index().sort_values('REGISTROS',ascending=False)
        hor.to_excel(writer,index=False,sheet_name='Horarios')
        loc=df.groupby('LOCALIDAD',dropna=False).size().reset_index(name='REGISTROS').sort_values('REGISTROS',ascending=False); loc.to_excel(writer,index=False,sheet_name='Localidades')
        if validation is not None and not validation.empty: validation.to_excel(writer,index=False,sheet_name='Control menciones')
    bio.seek(0); wb=load_workbook(bio)
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='17365D'); c.alignment=Alignment(horizontal='center',vertical='center')
        for col in ws.columns:
            letter=get_column_letter(col[0].column); ws.column_dimensions[letter].width=min(max(max(len(str(c.value or '')) for c in col)+2,10),55)
        if ws.title=='COI completo': ws.column_dimensions['L'].width=70
    out=io.BytesIO(); wb.save(out); return out.getvalue()

def pdf_pages(data,pages):
    src=fitz.open(stream=data,filetype='pdf'); out=fitz.open()
    for p in sorted(set(int(x) for x in pages)): out.insert_pdf(src,from_page=p-1,to_page=p-1)
    b=out.tobytes(); out.close(); src.close(); return b
